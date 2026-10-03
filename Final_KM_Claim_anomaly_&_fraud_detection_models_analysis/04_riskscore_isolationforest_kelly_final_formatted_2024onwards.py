"""
Step 4 - Kelly-aligned Isolation Forest scoring and revised-style Excel output.

This version keeps the Kelly-approved scoring logic, but writes the final workbook
in the same business-facing structure as 04_claims_riskscored_revised.xlsx.

Workbook order
--------------
1. read_me
2. why_each_claim
3. claims_individual
4. by_claimant
5. by_dealer
6. final_features_step4
7. features_used_in_model
8. features_excluded
9. missing_value_handling
10. leakage_diagnostics
11. model_settings

Important controls
------------------
* Only Processed and Rejected records are scored.
* Current claim outcomes and business-rule fields are excluded from Isolation Forest.
* Historical outcome-rate fields are retained for audit, but excluded from the model.
* Missing numeric inputs are median-imputed and documented.
* why_each_claim begins with the same A:R columns as the revised workbook.
* claims_individual begins with the same original business columns as the revised workbook.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from sklearn.ensemble import IsolationForest

BASE = Path(__file__).resolve().parent
DATE_CUTOFF = pd.Timestamp("2024-01-01")
INPUT_FILE = BASE / "data" / "03_claims_features_kelly_final.parquet"
OUTPUT_STUB = "04_claims_riskscored_kelly_final_2024onwards"

CONTAMINATION = 0.02
HIGH_PERCENTILE = 90
MEDIUM_PERCENTILE = 75
WILSON_Z = 1.96
MEDIUM_LIFT = 1.5
MEDIUM_EB_PROB = 0.80

MODEL_FEATURES = [
    "is_serialized_product", "claim_days_since_sale", "claim_file_delay", "claim_month",
    "product_seasonal_month_claims", "product_seasonal_avg_month_claims",
    "product_seasonal_active_months", "product_seasonality_claim_count_ratio",
    "product_seasonality_claim_count_high_flag", "claim_total_qty",
    "claim_claimant_recency", "claim_dealer_recency", "claimant_avg_frequency",
    "claimant_claim_count", "claimant_avg_total_qty", "z_claimant_avg_total_qty",
    "dealer_avg_total_qty", "z_dealer_avg_total_qty",
    "claimant_recent_claim_frequency", "claim_product_recency", "prod_avg_frequency",
    "prod_avg_qty_claim", "quantity_z_score", "claimant_avg_prod_qty",
    "z_claimant_avg_prod_qty", "r_prod_avg_qty_claim",
    "product_seasonality_quantity_ratio", "seasonality_impact_score",
    "claim_qty_month_comp", "claim_prod_amt", "claim_total_amt",
    "claimant_avg_total_amt", "dealer_avg_total_amt", "prod_avg_amt_claim",
    "claimant_avg_prod_amt", "r_prod_avg_amt_claim", "amount_z_score",
    "z_claimant_avg_total_amt", "z_dealer_avg_total_amt",
    "z_claimant_avg_prod_amt", "missing_claim_submission_date_flag",
    "missing_product_sold_date_flag", "claim_prod_qty", "claim_qty_month_z",
]

# Historical workflow outcomes remain visible for audit but are not model inputs.
HISTORICAL_OUTCOME_FEATURES = {
    "claimant_reject_rate", "claimant_return_rate", "claimant_cancel_rate",
    "claimant_percent_returned_index", "dealer_reject_rate",
    "dealer_return_rate", "dealer_cancel_rate",
}

PROHIBITED_MODEL_COLUMNS = {
    "ClaimId", "ClaimFormID", "ClaimItemId", "ClaimantName", "DealerName",
    "ClaimItemStatus", "ClaimItemStatusReason", "returned_claim", "cancelled_claim",
    "rejected_claim", "processed_claim", "rejected_claim_event", "returned_claim_event",
    "cancelled_claim_event", "processed_claim_event", "scoring_eligible",
    "anomaly_behavior_event", "history_anomaly_behavior_indicator",
    "history_anomaly_behavior_count", "history_anomaly_behavior_rate", "anamoly_flag",
    "fraud_flag", "failed_validation", "num_days", "ItemQuantity", "SaleQuantity",
    "IsSerializedProduct", "claim_days_since_sale_raw", "negative_claim_timing_flag",
    "date_eligible",
} | HISTORICAL_OUTCOME_FEATURES

WHY_EACH_CLAIM_FRONT_COLUMNS = [
    "ClaimantName", "ClaimFormID", "ClaimId", "DealerName", "ProductID",
    "ClaimItemStatus", "risk_level", "risk_score", "risk_score_1000",
    "risk_percentile", "risk_rank", "audit_band", "Why this rating",
    "Driver 1 (field = value)", "Driver 2 (field = value)",
    "Driver 3 (field = value)", "features_used_for_rating", "risk_reason",
]

# These are the first columns used by claims_individual in the revised workbook.
CLAIMS_INDIVIDUAL_FRONT_COLUMNS = [
    "ClaimFormID", "ClaimName", "ProductSoldDate", "ClaimSubmitionDate",
    "ClaimId", "ClaimItemId", "ItemDescription", "ItemQuantity",
    "ClaimItemStatusDate", "SerialNumber", "ClaimantName", "ClaimantName.1",
    "DealerName", "ProductID", "ProductStatus", "ProductName",
    "IsSerializedProduct", "ProductName.1",
]

RISK_ORDER = {"Low": 0, "Medium": 1, "High": 2}


def read_data(path) -> pd.DataFrame:
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = (BASE / p).resolve()
    if not p.exists():
        raise FileNotFoundError(f"Input file not found: {p}")
    if p.suffix.lower() == ".parquet":
        return pd.read_parquet(p)
    if p.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(p, engine="openpyxl")
    raise ValueError(f"Unsupported input type: {p.suffix}")


def select_model_features(df: pd.DataFrame):
    """Select only approved, present, numeric and nonconstant model features."""
    used, excluded = [], []
    for feature in MODEL_FEATURES:
        if feature in PROHIBITED_MODEL_COLUMNS:
            excluded.append((feature, "Excluded: workflow outcome, identifier, or duplicate"))
        elif feature not in df.columns:
            excluded.append((feature, "Approved Kelly feature missing from Step 3 input"))
        else:
            values = pd.to_numeric(df[feature], errors="coerce")
            if values.notna().sum() == 0:
                excluded.append((feature, "No usable numeric values"))
            elif values.nunique(dropna=True) <= 1:
                excluded.append((feature, "Constant in this run"))
            else:
                used.append(feature)

    # Report historical outcome fields independently even though they are not candidates.
    for feature in sorted(HISTORICAL_OUTCOME_FEATURES):
        if feature in df.columns:
            excluded.append((feature, "Audit context only: prior workflow outcome rate"))
    if not used:
        raise ValueError("No usable Kelly model features remain after validation.")
    return used, pd.DataFrame(excluded, columns=["feature", "why_excluded"])


def build_model_matrix(df: pd.DataFrame, features: list[str]):
    """Create a finite numeric matrix and document median imputation."""
    matrix = df[features].apply(pd.to_numeric, errors="coerce")
    matrix = matrix.replace([np.inf, -np.inf], np.nan)
    report = []
    for feature in features:
        missing_count = int(matrix[feature].isna().sum())
        fill_value = matrix[feature].median()
        method = "median"
        if pd.isna(fill_value):
            fill_value, method = 0.0, "zero because entire feature is missing"
        matrix[feature] = matrix[feature].fillna(float(fill_value))
        report.append({
            "feature": feature,
            "missing_values_imputed": missing_count,
            "imputation_method": method,
            "imputation_value": float(fill_value),
        })
    if not np.isfinite(matrix.to_numpy(dtype=float)).all():
        raise ValueError("Non-finite values remain after missing-value handling.")
    return matrix.astype(float), pd.DataFrame(report)


def score_isolation_forest(matrix: pd.DataFrame, contamination: float):
    model = IsolationForest(
        n_estimators=200,
        contamination=contamination,
        random_state=42,
        n_jobs=-1,
    )
    model.fit(matrix)
    raw = -model.score_samples(matrix)
    low, high = float(raw.min()), float(raw.max())
    risk = 100.0 * (raw - low) / (high - low) if high > low else np.zeros_like(raw)
    return np.round(risk, 2), model


def assign_risk_levels(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    scores = pd.to_numeric(out["risk_score"], errors="coerce").fillna(0.0)
    high_cut = float(np.percentile(scores, HIGH_PERCENTILE))
    medium_cut = float(np.percentile(scores, MEDIUM_PERCENTILE))
    out["risk_level_model"] = np.where(
        scores >= high_cut, "High", np.where(scores >= medium_cut, "Medium", "Low")
    )
    # Business outcomes do not promote scores. Final level remains model-derived.
    out["risk_level"] = out["risk_level_model"]
    return out


def create_risk_reasons(df: pd.DataFrame) -> pd.Series:
    reasons = [[] for _ in range(len(df))]
    rules = [
        ("quantity_z_score", 3, "product quantity far above product norm"),
        ("amount_z_score", 3, "product amount far above product norm"),
        ("z_claimant_avg_total_qty", 3, "claim quantity far above claimant norm"),
        ("z_claimant_avg_total_amt", 3, "claim amount far above claimant norm"),
        ("z_dealer_avg_total_qty", 3, "claim quantity far above dealer norm"),
        ("z_dealer_avg_total_amt", 3, "claim amount far above dealer norm"),
        ("r_prod_avg_qty_claim", 3, "product quantity is 3x+ claimant-product norm"),
        ("r_prod_avg_amt_claim", 3, "product amount is 3x+ claimant-product norm"),
        ("claim_qty_month_z", 3, "quantity high after seasonal adjustment"),
    ]
    for column, threshold, phrase in rules:
        if column not in df.columns:
            continue
        hits = pd.to_numeric(df[column], errors="coerce").ge(threshold).fillna(False)
        for row_number in np.flatnonzero(hits.to_numpy()):
            reasons[row_number].append(phrase)
    return pd.Series(
        [", ".join(parts) if parts else "no single threshold driver; multivariate model score"
         for parts in reasons],
        index=df.index,
    )


def highest_level(values) -> str:
    levels = [str(value).strip().title() for value in values]
    levels = [level for level in levels if level in RISK_ORDER]
    return max(levels, key=RISK_ORDER.get) if levels else "Low"


def wilson_lower(k: int, n: int, z: float = WILSON_Z) -> float:
    if n <= 0:
        return 0.0
    p = k / n
    denominator = 1 + z * z / n
    center = p + z * z / (2 * n)
    margin = z * np.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return max(0.0, (center - margin) / denominator)


def build_entity_summary(df: pd.DataFrame, entity_column: str) -> pd.DataFrame:
    """Build one claimant/dealer row using unique ClaimId units where available."""
    required = [entity_column, "risk_score", "risk_level", "risk_level_model", "risk_reason"]
    missing = [column for column in required if column not in df.columns]
    if missing:
        raise ValueError(f"Cannot build {entity_column} summary; missing columns: {missing}")

    columns = required + (["ClaimId"] if "ClaimId" in df.columns else [])
    work = df[columns].copy()
    work[entity_column] = work[entity_column].astype("string").fillna("[Missing]").str.strip()
    if "ClaimId" in work.columns and work["ClaimId"].notna().any():
        key = work["ClaimId"].astype("string")
        missing_id = key.isna() | key.str.strip().eq("")
        row_keys = pd.Series([f"__ROW_{i}" for i in range(len(work))], index=work.index)
        work["_claim_key"] = key.mask(missing_id, row_keys)
        unit = "unique ClaimId within entity"
    else:
        work["_claim_key"] = [f"__ROW_{i}" for i in range(len(work))]
        unit = "input row"

    base = work.groupby([entity_column, "_claim_key"], dropna=False).agg(
        risk_score=("risk_score", "max"),
        risk_level=("risk_level", highest_level),
        risk_level_model=("risk_level_model", highest_level),
        risk_reason=("risk_reason", lambda s: "; ".join(dict.fromkeys(str(x) for x in s))),
    ).reset_index()

    baseline = float(base["risk_level"].eq("High").mean())
    summary = base.assign(_high=base["risk_level"].eq("High").astype(int)).groupby(
        entity_column, dropna=False
    ).agg(
        claim_count=("_high", "size"),
        high_count=("_high", "sum"),
        avg_risk_score=("risk_score", "mean"),
        max_risk_score=("risk_score", "max"),
    ).reset_index()
    summary["raw_high_rate"] = summary["high_count"] / summary["claim_count"]
    summary["wilson_lower"] = [wilson_lower(int(k), int(n)) for k, n in zip(
        summary["high_count"], summary["claim_count"])]
    summary["wilson_lift"] = summary["wilson_lower"] / baseline if baseline > 0 else 0.0

    # Conservative empirical-Bayes prior centered on the observed population rate.
    prior_strength = 100.0
    alpha = max(baseline * prior_strength, 1e-6)
    beta = max((1 - baseline) * prior_strength, 1e-6)
    summary["eb_rate"] = (summary["high_count"] + alpha) / (
        summary["claim_count"] + alpha + beta)
    if 0 < baseline < 1:
        summary["eb_prob_above_baseline"] = [
            1 - stats.beta.cdf(baseline, k + alpha, n - k + beta)
            for k, n in zip(summary["high_count"], summary["claim_count"])
        ]
        summary["binom_p"] = [
            stats.binomtest(int(k), int(n), baseline, alternative="greater").pvalue
            for k, n in zip(summary["high_count"], summary["claim_count"])
        ]
    else:
        summary["eb_prob_above_baseline"] = 0.0
        summary["binom_p"] = 1.0

    high = summary["wilson_lower"] > baseline
    medium = (~high) & (baseline > 0) & (
        summary["eb_rate"] >= MEDIUM_LIFT * baseline
    ) & (summary["eb_prob_above_baseline"] >= MEDIUM_EB_PROB)
    prefix = "claimant" if entity_column == "ClaimantName" else "dealer"
    summary[f"{prefix}_risk_level"] = np.where(high, "High", np.where(medium, "Medium", "Low"))
    summary[f"{prefix}_risk_level_model"] = summary[f"{prefix}_risk_level"]
    summary[f"{prefix}_risk_reason"] = np.where(
        high, "High-claim rate is above the population baseline after Wilson adjustment",
        np.where(medium, "Elevated Empirical-Bayes rate", "No statistically elevated High-claim rate"),
    )
    summary["population_high_rate"] = baseline
    summary["counting_unit"] = unit

    front = [entity_column, f"{prefix}_risk_level_model", f"{prefix}_risk_level",
             f"{prefix}_risk_reason"]
    remaining = [column for column in summary.columns if column not in front]
    return summary[front + remaining].sort_values(
        [f"{prefix}_risk_level", "wilson_lift"],
        key=lambda s: s.map({"High": 0, "Medium": 1, "Low": 2}) if s.name.endswith("risk_level") else s,
        ascending=[True, False],
    ).reset_index(drop=True)


def ensure_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    out = df.copy()
    for column in columns:
        if column not in out.columns:
            out[column] = pd.NA
    return out


def build_why_each_claim(df: pd.DataFrame, matrix: pd.DataFrame, used_features: list[str]):
    """Create the revised workbook's A:R explanation layout, then append all remaining columns."""
    view = df.copy()
    score = pd.to_numeric(view["risk_score"], errors="coerce").fillna(0.0)
    view["risk_score_1000"] = (score * 10).round(0).astype("Int64")
    view["risk_percentile"] = score.rank(method="average", pct=True).mul(100).round(2)
    view["risk_rank"] = score.rank(method="min", ascending=False).astype("Int64")
    view["audit_band"] = view["risk_level"].map({
        "High": "Priority audit",
        "Medium": "Review if capacity allows",
        "Low": "Routine monitoring",
    }).fillna("Routine monitoring")

    numeric_matrix = matrix[used_features].astype(float)
    median = numeric_matrix.median(axis=0)
    mad = (numeric_matrix - median).abs().median(axis=0).replace(0, np.nan)
    strength = ((numeric_matrix - median).abs().div(mad)).replace(
        [np.inf, -np.inf], np.nan).fillna(0.0)
    top_positions = np.argsort(-strength.to_numpy(), axis=1)[:, :min(3, len(used_features))]
    driver_columns = ["Driver 1 (field = value)", "Driver 2 (field = value)",
                      "Driver 3 (field = value)"]
    for driver_number, output_column in enumerate(driver_columns):
        values = []
        for row_number, positions in enumerate(top_positions):
            if driver_number >= len(positions):
                values.append("")
                continue
            feature = used_features[int(positions[driver_number])]
            original_value = view.iloc[row_number][feature] if feature in view.columns else np.nan
            values.append(f"{feature} = {original_value}")
        view[output_column] = values

    view["features_used_for_rating"] = ", ".join(used_features)
    business_reason = view.get(
        "business_review_reason",
        pd.Series("No business rule triggered", index=view.index),
    )
    view["Why this rating"] = (
        "Model level: " + view["risk_level"].astype("string")
        + "; model reason: " + view["risk_reason"].astype("string")
        + "; business review: " + business_reason.astype("string")
    )
    view = ensure_columns(view, WHY_EACH_CLAIM_FRONT_COLUMNS)
    remaining = [column for column in view.columns if column not in WHY_EACH_CLAIM_FRONT_COLUMNS]
    return view[WHY_EACH_CLAIM_FRONT_COLUMNS + remaining]


def build_claims_individual(df: pd.DataFrame) -> pd.DataFrame:
    """Place the revised workbook's original business columns first, then all remaining fields."""
    view = ensure_columns(df, CLAIMS_INDIVIDUAL_FRONT_COLUMNS)
    remaining = [column for column in view.columns if column not in CLAIMS_INDIVIDUAL_FRONT_COLUMNS]
    return view[CLAIMS_INDIVIDUAL_FRONT_COLUMNS + remaining]


def build_read_me() -> pd.DataFrame:
    return pd.DataFrame([
        {"section": "Workbook purpose", "explanation": "Shows model anomaly results, claimant/dealer summaries, and separate business-defined review signals."},
        {"section": "Date cutoff", "explanation": "Only claims with ClaimSubmitionDate on or after January 1, 2024 are included in scoring and entity summaries."},
        {"section": "why_each_claim", "explanation": "Starts with the revised A:R explanation layout, followed by all remaining claim fields."},
        {"section": "claims_individual", "explanation": "One row per scored input record, with original business columns first and all engineered/scoring fields afterward."},
        {"section": "by_claimant", "explanation": "One consolidated claimant view using unique ClaimId units, Wilson adjustment, and Empirical Bayes context."},
        {"section": "by_dealer", "explanation": "One consolidated dealer view using the same statistical method."},
        {"section": "features_used_in_model", "explanation": "Exact numeric features sent to Isolation Forest."},
        {"section": "features_excluded", "explanation": "Features not used and the reason for exclusion."},
        {"section": "missing_value_handling", "explanation": "Documents median imputation for each model input."},
        {"section": "leakage_diagnostics", "explanation": "Checks statuses, prohibited fields, business rules, and matrix completeness."},
        {"section": "Important separation", "explanation": "Current claim outcomes, historical outcome rates, history_anomaly_behavior_indicator, and anamoly_flag are visible for audit but do not affect Isolation Forest scores."},
    ])


def build_feature_catalogue(df: pd.DataFrame, used_features: list[str]) -> pd.DataFrame:
    candidates = list(dict.fromkeys(MODEL_FEATURES + sorted(HISTORICAL_OUTCOME_FEATURES)))
    rows = []
    for feature in candidates:
        if feature in HISTORICAL_OUTCOME_FEATURES:
            role = "Audit context only: historical workflow outcome rate"
        elif feature in used_features:
            role = "Used by Isolation Forest"
        elif feature not in df.columns:
            role = "Missing from Step 3 input"
        else:
            role = "Excluded because missing/constant/unusable"
        rows.append({
            "feature": feature,
            "present_in_step3": feature in df.columns,
            "used_by_isolation_forest": feature in used_features,
            "missing_pct_before_imputation": (
                round(float(df[feature].isna().mean()), 6) if feature in df.columns else np.nan
            ),
            "role": role,
        })
    return pd.DataFrame(rows)


def apply_revised_formatting(path: Path):
    """Apply the revised workbook's readable styles without changing data."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    dark_blue = PatternFill("solid", fgColor="1F4E78")
    white_bold = Font(color="FFFFFF", bold=True)
    risk_fills = {
        "High": PatternFill("solid", fgColor="F4CCCC"),
        "Medium": PatternFill("solid", fgColor="FFF2CC"),
        "Low": PatternFill("solid", fgColor="D9EAD3"),
    }

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = dark_blue
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 30

        # Use a sample for widths so very large workbooks remain practical.
        sample_rows = min(worksheet.max_row, 400)
        for column_number in range(1, worksheet.max_column + 1):
            header = str(worksheet.cell(1, column_number).value or "")
            max_length = len(header)
            for row_number in range(2, sample_rows + 1):
                value = worksheet.cell(row_number, column_number).value
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            worksheet.column_dimensions[get_column_letter(column_number)].width = min(max_length + 2, 45)

    # Revised workbook behavior: keep explanation columns visible while scrolling.
    if "why_each_claim" in workbook.sheetnames:
        worksheet = workbook["why_each_claim"]
        worksheet.freeze_panes = "G2"
        headers = {str(cell.value): cell.column for cell in worksheet[1]}
        for header in ["Why this rating", "Driver 1 (field = value)",
                       "Driver 2 (field = value)", "Driver 3 (field = value)",
                       "features_used_for_rating", "risk_reason"]:
            if header in headers:
                worksheet.column_dimensions[get_column_letter(headers[header])].width = 45

    for sheet_name in ["why_each_claim", "claims_individual", "by_claimant", "by_dealer"]:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        headers = {str(cell.value): cell.column for cell in worksheet[1]}
        risk_header = next((name for name in ["risk_level", "claimant_risk_level", "dealer_risk_level"]
                            if name in headers), None)
        if risk_header:
            column_number = headers[risk_header]
            for row_number in range(2, worksheet.max_row + 1):
                level = str(worksheet.cell(row_number, column_number).value or "").title()
                fill = risk_fills.get(level)
                if fill:
                    for cell in worksheet[row_number]:
                        cell.fill = fill

        for percentage_column in ["raw_high_rate", "wilson_lower", "eb_rate",
                                  "eb_prob_above_baseline", "population_high_rate"]:
            if percentage_column in headers:
                letter = get_column_letter(headers[percentage_column])
                for cell in worksheet[letter][1:]:
                    cell.number_format = "0.0%"

        for score_column in ["risk_score", "avg_risk_score", "max_risk_score",
                             "risk_score_1000", "risk_percentile"]:
            if score_column in headers:
                letter = get_column_letter(headers[score_column])
                for cell in worksheet[letter][1:]:
                    cell.number_format = "0.00"

    workbook.save(path)


def run(input_path=INPUT_FILE, output_stub=OUTPUT_STUB, contamination=CONTAMINATION):
    df = read_data(input_path)
    if df.empty:
        raise ValueError("Step 4 input is empty.")

    # Defensive enforcement: score only claims submitted on or after 2024-01-01.
    df["ClaimSubmitionDate"] = pd.to_datetime(df["ClaimSubmitionDate"], errors="coerce")
    rows_before_cutoff = len(df)
    df = df.loc[df["ClaimSubmitionDate"].ge(DATE_CUTOFF)].copy().reset_index(drop=True)
    df["date_eligible"] = 1
    print(f"[Step 4] 2024 cutoff: kept {len(df):,} of {rows_before_cutoff:,} rows")
    if df.empty: raise ValueError("No claims remain after the January 1, 2024 cutoff.")

    status = df["ClaimItemStatus"].astype("string").str.strip().str.casefold()
    df = df.loc[status.isin(["processed", "rejected"])].copy().reset_index(drop=True)
    if df.empty:
        raise ValueError("No Processed or Rejected claims remain for scoring.")

    used_features, excluded_features = select_model_features(df)
    prohibited_used = sorted(set(used_features) & PROHIBITED_MODEL_COLUMNS)
    if prohibited_used:
        raise ValueError(f"Prohibited leakage fields selected: {prohibited_used}")

    matrix, imputation_report = build_model_matrix(df, used_features)
    df["risk_score"], _ = score_isolation_forest(matrix, contamination)
    df = assign_risk_levels(df)
    df["risk_reason"] = create_risk_reasons(df)

    for column in ["history_anomaly_behavior_indicator", "anamoly_flag",
                   "negative_claim_timing_flag"]:
        if column not in df.columns:
            df[column] = 0
    df["business_review_flag"] = (
        df["history_anomaly_behavior_indicator"].eq(1)
        | df["anamoly_flag"].eq(1)
        | df["negative_claim_timing_flag"].eq(1)
    ).astype("int8")
    df["business_review_reason"] = np.select(
        [
            df["negative_claim_timing_flag"].eq(1),
            df["history_anomaly_behavior_indicator"].eq(1) & df["anamoly_flag"].eq(1),
            df["history_anomaly_behavior_indicator"].eq(1),
            df["anamoly_flag"].eq(1),
        ],
        [
            "Negative claim timing requires data review",
            "Prior anomaly behavior and 90+ day filing delay",
            "Prior anomaly behavior",
            "90+ day filing delay",
        ],
        default="No business rule triggered",
    )

    why_each_claim = build_why_each_claim(df, matrix, used_features)
    claims_individual = build_claims_individual(df)
    by_claimant = build_entity_summary(df, "ClaimantName")
    by_dealer = build_entity_summary(df, "DealerName")
    feature_catalogue = build_feature_catalogue(df, used_features)

    diagnostics = pd.DataFrame([
        {"test": "only_processed_rejected_scored", "passed": bool(df["ClaimItemStatus"].astype("string").str.casefold().isin(["processed", "rejected"]).all()), "detail": "Scoring population"},
        {"test": "only_2024_and_later_scored", "passed": bool(pd.to_datetime(df["ClaimSubmitionDate"], errors="coerce").ge(DATE_CUTOFF).all()), "detail": "ClaimSubmitionDate is January 1, 2024 or later"},
        {"test": "all_used_features_approved", "passed": set(used_features).issubset(MODEL_FEATURES), "detail": f"{len(used_features)} features used"},
        {"test": "no_prohibited_columns_in_model", "passed": not prohibited_used, "detail": ", ".join(prohibited_used) or "None"},
        {"test": "historical_outcome_rates_excluded", "passed": not (set(used_features) & HISTORICAL_OUTCOME_FEATURES), "detail": "Historical workflow rates are audit context only"},
        {"test": "finite_model_matrix", "passed": bool(np.isfinite(matrix.to_numpy()).all()), "detail": "After documented median imputation"},
        {"test": "why_each_claim_has_A_to_R_layout", "passed": list(why_each_claim.columns[:18]) == WHY_EACH_CLAIM_FRONT_COLUMNS, "detail": "Revised explanation layout"},
    ])
    if not diagnostics["passed"].all():
        raise ValueError("Step 4 validation failed; review leakage_diagnostics.")

    read_me = build_read_me()
    settings = pd.DataFrame([
        {"setting": "input_rows_scored", "value": len(df), "meaning": "Processed and Rejected input rows"},
        {"setting": "date_cutoff", "value": "2024-01-01", "meaning": "Only claims submitted on or after this date are scored"},
        {"setting": "model_feature_count", "value": len(used_features), "meaning": "Approved, present, numeric, nonconstant features"},
        {"setting": "contamination", "value": contamination, "meaning": "Isolation Forest fit setting"},
        {"setting": "high_percentile", "value": HIGH_PERCENTILE, "meaning": "Global score threshold for High"},
        {"setting": "medium_percentile", "value": MEDIUM_PERCENTILE, "meaning": "Global score threshold for Medium"},
        {"setting": "entity_counting_unit", "value": "unique ClaimId within entity when available", "meaning": "Prevents repeated claim-item rows from inflating entity statistics"},
    ])
    features_used = pd.DataFrame({
        "model_feature": used_features,
        "dtype": [str(df[column].dtype) for column in used_features],
        "missing_pct": [round(float(df[column].isna().mean()), 6) for column in used_features],
        "unique_values": [int(df[column].nunique(dropna=True)) for column in used_features],
        "role": "Used by Isolation Forest",
    })

    data_dir = BASE / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    parquet_path = data_dir / f"{output_stub}.parquet"
    excel_path = data_dir / f"{output_stub}.xlsx"
    df.to_parquet(parquet_path, index=False)

    # Sheet write order is deliberate and matches the requested business layout.
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        read_me.to_excel(writer, sheet_name="read_me", index=False)
        why_each_claim.to_excel(writer, sheet_name="why_each_claim", index=False)
        claims_individual.to_excel(writer, sheet_name="claims_individual", index=False)
        by_claimant.to_excel(writer, sheet_name="by_claimant", index=False)
        by_dealer.to_excel(writer, sheet_name="by_dealer", index=False)
        feature_catalogue.to_excel(writer, sheet_name="final_features_step4", index=False)
        features_used.to_excel(writer, sheet_name="features_used_in_model", index=False)
        excluded_features.to_excel(writer, sheet_name="features_excluded", index=False)
        imputation_report.to_excel(writer, sheet_name="missing_value_handling", index=False)
        diagnostics.to_excel(writer, sheet_name="leakage_diagnostics", index=False)
        settings.to_excel(writer, sheet_name="model_settings", index=False)

    apply_revised_formatting(excel_path)
    print(f"[save] {parquet_path}\n[save] {excel_path}")
    return df, by_claimant, by_dealer


def main():
    parser = argparse.ArgumentParser(
        description="Kelly Step 4 Isolation Forest scoring with revised-style workbook output"
    )
    parser.add_argument("--input", default=INPUT_FILE)
    parser.add_argument("--output-stub", default=OUTPUT_STUB)
    parser.add_argument("--contamination", type=float, default=CONTAMINATION)
    args = parser.parse_args()
    run(args.input, args.output_stub, args.contamination)


if __name__ == "__main__":
    main()
