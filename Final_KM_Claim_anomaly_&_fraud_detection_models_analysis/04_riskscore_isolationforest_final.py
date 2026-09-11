"""
04_riskscore_isolationforest.py
Isolation Forest claim scoring plus claimant/dealer consolidated ratings

WHAT THIS STEP DOES
-------------------
Reads engineered claim features and produces:

1. claims_individual
   One row per INPUT RECORD, with:
   risk_score, risk_level_model, risk_level, and risk_reason.

2. by_claimant
   One consolidated rating per claimant.

3. by_dealer
   One consolidated rating per dealer.

IMPORTANT STATISTICAL NOTES
---------------------------
* HIGH_PERCENTILE=90 makes approximately the top 10% of ALL claim-record scores
  model-High before business-rule promotions. It does not force 10% High within
  every claimant or dealer.

* CONTAMINATION is passed to Isolation Forest, but exported risk tiers are based
  on global score percentiles. Therefore contamination does not directly set the
  percentage of exported High claims.

* Business-rule promotions can increase the final High rate above 10%.

* Entity consolidation uses one record per unique ClaimId within each entity
  when ClaimId is available. If ClaimId is unavailable, each input row is used.
  For duplicate rows belonging to the same entity and ClaimId, the maximum risk
  score and most severe risk level are retained for entity statistics.

ENTITY CONSOLIDATION
--------------------
An entity is not rated from the count of High claims alone. Its High-claim rate
is compared with the population High-claim rate using:

* Wilson lower bound: primary High decision, conservative and sample-size-aware.
* Empirical Bayes: stabilizes small entities by shrinking rates toward baseline.
* Exact binomial test: supporting evidence reported as a one-sided p-value.

Final entity decision:
* High: Wilson lower bound is above population High rate.
* Medium: not High, EB rate is at least MEDIUM_LIFT times baseline, and the
  posterior probability of being above baseline is at least MEDIUM_EB_PROB.
* Low: otherwise.

MEDIUM_LIFT and MEDIUM_EB_PROB are configurable audit-policy thresholds, not
universal statistical constants.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from sklearn.ensemble import IsolationForest


BASE = Path(__file__).resolve().parent
INPUT_FILE = BASE / "data" / "03_claims_features.parquet"
PROFILE_FILE = BASE / "data" / "03b_feature_profile.xlsx"
OUTPUT_STUB = "04_claims_riskscored"


#OUTPUT_FILE = BASE / "data" / "04_claims_riskscored.parquet"
#EXCEL_OUTPUT = BASE / "data" / "04_claims_riskscored.xlsx"


# Claim-level scoring settings.
CONTAMINATION = 0.02
HIGH_PERCENTILE = 90
MEDIUM_PERCENTILE = 75

# Entity-level settings.
WILSON_Z = 1.96
MEDIUM_LIFT = 1.5
MEDIUM_EB_PROB = 0.80
MIN_PRIOR_STRENGTH = 2.0
MAX_PRIOR_STRENGTH = 10_000.0
FALLBACK_PRIOR_STRENGTH = 100.0

# Used only when the profile workbook is unavailable.
MODEL_FEATURES = [
    "z_claimant_avg_total_qty", "z_claimant_avg_total_amt",
    "z_claimant_avg_prod_qty", "z_claimant_avg_prod_amt",
    "z_dealer_avg_total_qty", "z_dealer_avg_total_amt",
    "quantity_z_score", "amount_z_score",
    "rolling_30d_claim_count_z", "claim_qty_month_z",
    "recency_vs_avg_frequency", "r_prod_avg_qty_claim", "r_prod_avg_amt_claim",
    "claimant_recent_claim_frequency", "claimant_percent_returned_index",
    "product_seasonality_quantity_ratio", "seasonality_impact_score",
    "claimant_reject_rate", "claimant_return_rate", "claimant_cancel_rate",
    "claimant_failed_validation_rate",
    "dealer_reject_rate", "dealer_return_rate", "dealer_cancel_rate",
    "dealer_failed_validation_rate",
    "claim_days_since_sale", "claim_file_delay", "rejected_claim",
    "failed_validation",
]

ALWAYS_EXCLUDE = [
    "ClaimFormID", "ClaimId", "ClaimItemId", "ClaimantName", "DealerName",
    "ProductID", "ProductName", "ProductDescription", "SerialNumber", "ClaimName",
    "SalesTransactionID", "AuditComment", "ClaimItemStatus", "ClaimItemStatusReason",
    "IsSerializedProduct", "ProductStatus", "ProductManafacturer",
    "ClaimSubmitionDate", "ProductSoldDate", "ClaimDate", "InvoiceDate",
    "ClaimItemStatusDate", "product_claim_month", "claim_month",
    "risk_score", "risk_level", "risk_level_model", "risk_reason",
]

LEVEL_ORDER = {"Low": 0, "Medium": 1, "High": 2}


def pv(value):
    """Resolve a path relative to this script unless it is already absolute."""
    p = Path(value).expanduser()
    return (p if p.is_absolute() else BASE / p).resolve()


def read(path):
    """Read parquet or Excel input."""
    p = pv(path)
    if not p.exists():
        raise FileNotFoundError(f"Input file not found: {p}")
    if p.suffix.lower() == ".parquet":
        return pd.read_parquet(p)
    if p.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(p, engine="openpyxl")
    raise ValueError(f"Unsupported input type: {p.suffix}")


# ---------------------------------------------------------------------------
# Feature selection
# ---------------------------------------------------------------------------

def select_features_from_profile(df):
    """Select approved, usable numeric features and explain exclusions."""
    drop_rows = []

    if PROFILE_FILE.exists():
        profile = pd.read_excel(PROFILE_FILE, engine="openpyxl")
        required = {"column"}
        missing = required - set(profile.columns)
        if missing:
            raise ValueError(
                f"Profile is missing required column(s): {sorted(missing)}"
            )

        if "use_in_model" not in profile.columns:
            def derive(row):
                action = str(row.get("recommended_action", "")).lower()
                dtype = str(row.get("dtype", "")).lower()
                numeric = dtype.startswith(("int", "float"))
                return int(action.startswith("keep") and numeric)

            profile["use_in_model"] = profile.apply(derive, axis=1)

        use_flag = pd.to_numeric(profile["use_in_model"], errors="coerce").fillna(0)
        approved = profile.loc[use_flag == 1, "column"].astype(str).tolist()
        approved_set = set(approved)
        source = "profile (03b_feature_profile.xlsx)"

        for c in df.columns:
            if c in ALWAYS_EXCLUDE:
                drop_rows.append((c, "excluded: identifier/date/output"))
            elif c not in approved_set:
                match = profile.loc[profile["column"].astype(str) == str(c)]
                reason = ""
                if "recommended_action" in match.columns and not match.empty:
                    reason = str(match["recommended_action"].iloc[0])
                drop_rows.append((c, reason or "not approved by profile"))
    else:
        approved = MODEL_FEATURES
        source = "built-in MODEL_FEATURES (profile file not found)"

    final = []
    seen = set()

    for c in approved:
        if c in seen:
            continue
        seen.add(c)

        if c not in df.columns:
            drop_rows.append((c, "approved feature missing from input"))
            continue
        if c in ALWAYS_EXCLUDE:
            drop_rows.append((c, "excluded: identifier/date/output"))
            continue

        converted = pd.to_numeric(df[c], errors="coerce")
        numeric_count = int(converted.notna().sum())
        unique_count = int(converted.nunique(dropna=True))

        if numeric_count == 0:
            drop_rows.append((c, "no usable numeric values"))
            continue
        if unique_count <= 1:
            drop_rows.append((c, "constant or all-missing feature"))
            continue

        final.append(c)

    if not final:
        raise ValueError(
            "No usable model features remain. Check the profile and input columns."
        )

    dropped = pd.DataFrame(drop_rows, columns=["column", "why_dropped"])
    dropped = dropped.drop_duplicates().reset_index(drop=True)

    print(f"[features] source: {source}")
    print(f"[features] final model features: {len(final)}")
    return final, dropped


def build_model_matrix(df):
    """Build a finite numeric matrix for Isolation Forest."""
    features, dropped = select_features_from_profile(df)
    matrix = df[features].apply(pd.to_numeric, errors="coerce")
    matrix = matrix.replace([np.inf, -np.inf], np.nan).fillna(0.0)
    print(f"[info] model will use {len(features)} features on {len(matrix):,} rows")
    return matrix, features, dropped


# ---------------------------------------------------------------------------
# Isolation Forest score
# ---------------------------------------------------------------------------

def score_with_isolation_forest(matrix, contamination=CONTAMINATION, seed=42):
    """Fit Isolation Forest and min-max scale anomaly scores to 0-100."""
    if not 0 < contamination <= 0.5:
        raise ValueError("contamination must be greater than 0 and at most 0.5")

    model = IsolationForest(
        n_estimators=200,
        contamination=contamination,
        random_state=seed,
        n_jobs=-1,
    )
    model.fit(matrix)

    # Higher raw value means more anomalous.
    raw = -model.score_samples(matrix)
    lo, hi = float(raw.min()), float(raw.max())
    risk = 100.0 * (raw - lo) / (hi - lo) if hi > lo else np.zeros_like(raw)
    return np.round(risk, 2), model


# ---------------------------------------------------------------------------
# Per-record risk levels and reasons
# ---------------------------------------------------------------------------

def assign_risk_levels(df, score_col="risk_score"):
    """
    Turn each claim's 0-100 risk score into a High / Medium / Low label.

    Two steps:
      1. Rank by score: the top 10% become High, the next 15% Medium, rest Low.
      2. Safety rule: any claim that was rejected or failed validation is always
         marked High, even if its score was lower.
    """
    scores = pd.to_numeric(df[score_col], errors="coerce")

    # --- Step 1: label by score ranking ---
    # A claim in the top 10% of scores is "High", the next band is "Medium".
    high_cut = np.percentile(scores, HIGH_PERCENTILE)   # e.g. the 90th percentile
    med_cut  = np.percentile(scores, MEDIUM_PERCENTILE) # e.g. the 75th percentile

    df["risk_level_model"] = "Low"
    df.loc[scores >= med_cut,  "risk_level_model"] = "Medium"
    df.loc[scores >= high_cut, "risk_level_model"] = "High"

    # --- Step 2: business-rule safety net ---
    # Rejected or failed-validation claims always go to High for review.
    df["risk_level"] = df["risk_level_model"]
    for rule in ["rejected_claim", "failed_validation"]:
        if rule in df.columns:
            df.loc[df[rule] == 1, "risk_level"] = "High"

    print(f"[thresholds] Medium score >= {med_cut:.1f}, High score >= {high_cut:.1f}")
    return df








def build_reasons(df):
    """Create plain-language record-level reasons for strong input signals."""
    reasons = [[] for _ in range(len(df))]

    def flag(col, threshold, phrase, direction="high"):
        if col not in df.columns:
            return
        values = pd.to_numeric(df[col], errors="coerce").fillna(0)
        hit = values.ge(threshold) if direction == "high" else values.le(threshold)
        for i in np.flatnonzero(hit.to_numpy()):
            reasons[i].append(phrase)

    flag("quantity_z_score", 3, "product quantity far above product norm")
    flag("amount_z_score", 3, "product amount far above product norm")
    flag("z_claimant_avg_total_qty", 3, "claim quantity far above claimant norm")
    flag("z_claimant_avg_total_amt", 3, "claim amount far above claimant norm")
    flag("z_dealer_avg_total_qty", 3, "claim quantity far above dealer norm")
    flag("z_dealer_avg_total_amt", 3, "claim amount far above dealer norm")
    flag("r_prod_avg_qty_claim", 3, "product quantity 3x+ the claimant's usual")
    flag("r_prod_avg_amt_claim", 3, "product amount 3x+ the claimant's usual")
    flag(
        "recency_vs_avg_frequency",
        0.25,
        "claim filed unusually soon after the last",
        direction="low",
    )
    flag("rolling_30d_claim_count_z", 2, "unusual spike in recent claim activity")
    flag(
        "claimant_percent_returned_index",
        2,
        "claimant returned rate 2x+ the population",
    )
    flag("claim_qty_month_z", 3, "quantity high even after seasonal adjustment")
    flag("failed_validation", 1, "failed a validation rule")
    flag("rejected_claim", 1, "previously rejected by an auditor")
    flag("claim_file_delay", 1, "filed 90+ days after sale")

    return [
        ", ".join(items)
        if items
        else "no single strong signal (model-driven score)"
        for items in reasons
    ]


# ---------------------------------------------------------------------------
# Entity consolidation helpers
# ---------------------------------------------------------------------------

def wilson_lower(k, n, z=WILSON_Z):
    """Lower Wilson confidence bound for a binomial proportion."""
    if n <= 0:
        return 0.0
    p = k / n
    denominator = 1 + z * z / n
    center = p + z * z / (2 * n)
    margin = z * np.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return max(0.0, (center - margin) / denominator)


def highest_level(values):
    """Return the most severe normalized risk level in a group."""
    cleaned = [str(v).strip().title() for v in values]
    cleaned = [v for v in cleaned if v in LEVEL_ORDER]
    return max(cleaned, key=LEVEL_ORDER.get) if cleaned else "Low"


def combine_reasons(values):
    """Combine unique non-default reasons from duplicate claim records."""
    parts = []
    seen = set()
    for value in values:
        for part in str(value).split(","):
            text = part.strip()
            if not text or text.startswith("no single strong signal"):
                continue
            if text not in seen:
                seen.add(text)
                parts.append(text)
    return ", ".join(parts) if parts else "no single strong signal (model-driven score)"


def build_entity_claim_base(df, entity_col):
    """
    Build the unit used for entity statistics.
    Preferred unit: one unique ClaimId per entity.
    Fallback unit: one input row when ClaimId is unavailable.
    """
    needed = [entity_col, "risk_score", "risk_level", "risk_level_model", "risk_reason"]
    missing = [c for c in needed if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns for {entity_col} consolidation: {missing}")

    work = df[needed + (["ClaimId"] if "ClaimId" in df.columns else [])].copy()
    work[entity_col] = work[entity_col].astype("string").fillna("[Missing]").str.strip()
    work.loc[work[entity_col].eq(""), entity_col] = "[Missing]"

    if "ClaimId" in work.columns and work["ClaimId"].notna().any():
        # Missing IDs receive a unique row-level key so unrelated records are not merged.
        claim_key = work["ClaimId"].astype("string")
        missing_id = claim_key.isna() | claim_key.str.strip().eq("")
        row_keys = pd.Series(
            [f"__ROW_{i}" for i in range(len(work))], index=work.index, dtype="string"
        )
        work["_claim_key"] = claim_key.mask(missing_id, row_keys)
        unit_description = "unique ClaimId within entity"
    else:
        work["_claim_key"] = [f"__ROW_{i}" for i in range(len(work))]
        unit_description = "input row (ClaimId unavailable)"

    base = (
        work.groupby([entity_col, "_claim_key"], dropna=False)
        .agg(
            risk_score=("risk_score", "max"),
            risk_level=("risk_level", highest_level),
            risk_level_model=("risk_level_model", highest_level),
            risk_reason=("risk_reason", combine_reasons),
        )
        .reset_index()
    )

    print(
        f"[{entity_col}] entity unit: {unit_description}; "
        f"{len(work):,} input rows -> {len(base):,} entity-claim records"
    )
    return base


def fit_beta_prior(g, baseline):
    """
    Estimate a beta prior after removing approximate within-entity binomial noise.
    This is a method-of-moments empirical-Bayes estimate. It is more conservative
    than treating all observed variation in small-entity rates as real variation.
    """
    n = g["claim_count"].astype(float)
    rates = g["raw_high_rate"].astype(float)
    weights = n / n.sum()

    observed_var = float((weights * (rates - baseline) ** 2).sum())
    sampling_var = float((weights * baseline * (1 - baseline) / n).sum())
    between_var = max(observed_var - sampling_var, 0.0)

    if 0 < baseline < 1 and between_var > 0:
        strength = baseline * (1 - baseline) / between_var - 1
        strength = float(np.clip(strength, MIN_PRIOR_STRENGTH, MAX_PRIOR_STRENGTH))
    else:
        strength = FALLBACK_PRIOR_STRENGTH

    alpha = max(baseline * strength, 1e-6)
    beta = max((1 - baseline) * strength, 1e-6)
    return alpha, beta, observed_var, sampling_var, between_var


def calculate_entity_statistics(base, level_col):
    """Calculate Wilson, EB, and binomial statistics for one level definition."""
    levels = base[level_col].astype(str).str.strip().str.title()
    p0 = float(levels.eq("High").mean())

    temp = pd.DataFrame(
        {
            "entity": base.iloc[:, 0].values,
            "_high": levels.eq("High").astype(int).values,
            "_score": pd.to_numeric(base["risk_score"], errors="coerce").fillna(0).values,
        }
    )

    g = (
        temp.groupby("entity", dropna=False)
        .agg(
            claim_count=("_high", "size"),
            high_count=("_high", "sum"),
            avg_risk_score=("_score", "mean"),
            max_risk_score=("_score", "max"),
        )
        .reset_index()
    )
    g["raw_high_rate"] = g["high_count"] / g["claim_count"]
    g["wilson_lower"] = [
        wilson_lower(int(k), int(n))
        for k, n in zip(g["high_count"], g["claim_count"])
    ]
    g["wilson_lift"] = g["wilson_lower"] / p0 if p0 > 0 else 0.0

    alpha, beta, observed_var, sampling_var, between_var = fit_beta_prior(g, p0)
    g["eb_rate"] = (
        (g["high_count"] + alpha)
        / (g["claim_count"] + alpha + beta)
    )

    if 0 < p0 < 1:
        g["eb_prob_above_baseline"] = [
            1 - stats.beta.cdf(p0, k + alpha, n - k + beta)
            for k, n in zip(g["high_count"], g["claim_count"])
        ]
        g["binom_p"] = [
            stats.binomtest(int(k), int(n), p0, alternative="greater").pvalue
            for k, n in zip(g["high_count"], g["claim_count"])
        ]
    else:
        # No meaningful above-baseline test exists if every or no claim is High.
        g["eb_prob_above_baseline"] = 0.0
        g["binom_p"] = 1.0

    high = g["wilson_lower"] > p0
    medium = (
        (~high)
        & (p0 > 0)
        & (g["eb_rate"] >= MEDIUM_LIFT * p0)
        & (g["eb_prob_above_baseline"] >= MEDIUM_EB_PROB)
    )
    g["entity_risk_level"] = np.where(
        high, "High", np.where(medium, "Medium", "Low")
    )

    diagnostics = {
        "population_high_rate": p0,
        "prior_alpha": alpha,
        "prior_beta": beta,
        "observed_rate_variance": observed_var,
        "estimated_sampling_variance": sampling_var,
        "estimated_between_entity_variance": between_var,
    }
    return g, diagnostics


def summarize_entity_reasons(base, entity_col):
    """Summarize common reasons among each entity's High or Medium claim units."""
    temp = base[[entity_col, "risk_level", "risk_reason"]].copy()
    temp["_tier"] = temp["risk_level"].astype(str).str.strip().str.title()

    def summarize(sub):
        flagged = sub[sub["_tier"].isin(["High", "Medium"])]
        if flagged.empty:
            return "mostly low-risk claims"

        counts = {}
        for reason in flagged["risk_reason"].astype(str):
            for part in reason.split(","):
                text = part.strip()
                if text and not text.startswith("no single strong signal"):
                    counts[text] = counts.get(text, 0) + 1

        if not counts:
            return "elevated score without one dominant driver"

        top = sorted(counts.items(), key=lambda item: (-item[1], item[0]))[:4]
        return "; ".join(f"{reason} ({count} claims)" for reason, count in top)

    return (
        temp.groupby(entity_col, dropna=False)
        .apply(summarize, include_groups=False)
        .reset_index(name="entity_risk_reason")
    )


def consolidate_entity(df, entity_col):
    """Return one statistically adjusted row per claimant or dealer."""
    prefix = "claimant" if entity_col.lower().startswith("claimant") else "dealer"
    base = build_entity_claim_base(df, entity_col)

    final_stats, final_diag = calculate_entity_statistics(base, "risk_level")
    model_stats, model_diag = calculate_entity_statistics(base, "risk_level_model")

    final_stats = final_stats.rename(
        columns={
            "entity": entity_col,
            "entity_risk_level": f"{prefix}_risk_level",
        }
    )
    model_levels = model_stats[["entity", "entity_risk_level"]].rename(
        columns={
            "entity": entity_col,
            "entity_risk_level": f"{prefix}_risk_level_model",
        }
    )

    reasons = summarize_entity_reasons(base, entity_col).rename(
        columns={"entity_risk_reason": f"{prefix}_risk_reason"}
    )

    g = final_stats.merge(model_levels, on=entity_col, how="left")
    g = g.merge(reasons, on=entity_col, how="left")

    # Repeat audit-level assumptions/diagnostics on each row for transparent export.
    g["population_high_rate"] = final_diag["population_high_rate"]
    g["model_population_high_rate"] = model_diag["population_high_rate"]
    g["eb_prior_alpha"] = final_diag["prior_alpha"]
    g["eb_prior_beta"] = final_diag["prior_beta"]
    g["estimated_between_entity_variance"] = final_diag[
        "estimated_between_entity_variance"
    ]
    g["medium_lift_threshold"] = MEDIUM_LIFT
    g["medium_eb_probability_threshold"] = MEDIUM_EB_PROB
    g["counting_unit"] = (
        "unique ClaimId within entity"
        if "ClaimId" in df.columns and df["ClaimId"].notna().any()
        else "input row"
    )

    columns = [
        entity_col,
        f"{prefix}_risk_level_model",
        f"{prefix}_risk_level",
        f"{prefix}_risk_reason",
        "claim_count",
        "high_count",
        "raw_high_rate",
        "wilson_lower",
        "wilson_lift",
        "eb_rate",
        "eb_prob_above_baseline",
        "binom_p",
        "avg_risk_score",
        "max_risk_score",
        "population_high_rate",
        "model_population_high_rate",
        "eb_prior_alpha",
        "eb_prior_beta",
        "estimated_between_entity_variance",
        "medium_lift_threshold",
        "medium_eb_probability_threshold",
        "counting_unit",
    ]

    rank = {"High": 0, "Medium": 1, "Low": 2}
    g["_sort"] = g[f"{prefix}_risk_level"].map(rank).fillna(3)
    g = g.sort_values(
        ["_sort", "wilson_lift", "eb_rate"], ascending=[True, False, False]
    ).drop(columns="_sort")

    numeric_cols = g.select_dtypes(include=[np.number]).columns
    g[numeric_cols] = g[numeric_cols].round(4)
    return g[columns].reset_index(drop=True)


# ---------------------------------------------------------------------------
# Save and orchestrate
# ---------------------------------------------------------------------------

def make_excel_readable(path):
    """Apply simple Excel usability formatting without changing the data."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        sample_limit = min(ws.max_row, 500)
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(1, col_idx).value or "")
            max_len = len(header)
            for row_idx in range(2, sample_limit + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

        headers = {str(c.value): c.column for c in ws[1]}
        for pct_col in [
            "raw_high_rate", "wilson_lower", "eb_rate",
            "eb_prob_above_baseline", "population_high_rate",
            "model_population_high_rate", "medium_eb_probability_threshold",
        ]:
            if pct_col in headers:
                letter = get_column_letter(headers[pct_col])
                for cell in ws[letter][1:]:
                    cell.number_format = "0.0%"

        for score_col in ["risk_score", "avg_risk_score", "max_risk_score"]:
            if score_col in headers:
                letter = get_column_letter(headers[score_col])
                for cell in ws[letter][1:]:
                    cell.number_format = "0.00"

    wb.save(path)


def run(
    input_path=INPUT_FILE,
    output_stub=OUTPUT_STUB,
    contamination=CONTAMINATION,
    save_excel=True,
):
    """Run claim scoring, entity consolidation, and output generation."""
    df = read(input_path)
    if df.empty:
        raise ValueError("Input dataset is empty")

    matrix, used_features, dropped = build_model_matrix(df)

    if not dropped.empty:
        print("\n[dropped features - not sent to the model]")
        for _, row in dropped.iterrows():
            print(f"  {str(row['column']):<38} {row['why_dropped']}")
        print()

    # Stage 1: score every input record.
    df["risk_score"], _ = score_with_isolation_forest(matrix, contamination)
    df = assign_risk_levels(df)
    df["risk_reason"] = build_reasons(df)

    distribution = df["risk_level"].value_counts()
    total = len(df)
    print("[claim-record distribution]")
    for tier in ["High", "Medium", "Low"]:
        count = int(distribution.get(tier, 0))
        print(f"  {tier:6}: {count:,} ({100 * count / total:.1f}%)")

    # Stage 2: statistically adjusted entity ratings.
    claimant = (
        consolidate_entity(df, "ClaimantName")
        if "ClaimantName" in df.columns
        else None
    )
    dealer = (
        consolidate_entity(df, "DealerName")
        if "DealerName" in df.columns
        else None
    )

    if claimant is not None:
        counts = claimant["claimant_risk_level"].value_counts()
        print(
            "[claimant ratings] "
            + ", ".join(
                f"{tier} {int(counts.get(tier, 0))}"
                for tier in ["High", "Medium", "Low"]
            )
        )

    if dealer is not None:
        counts = dealer["dealer_risk_level"].value_counts()
        print(
            "[dealer ratings]   "
            + ", ".join(
                f"{tier} {int(counts.get(tier, 0))}"
                for tier in ["High", "Medium", "Low"]
            )
        )

    # Stage 3: save outputs.
    data_dir = BASE / "data"
    data_dir.mkdir(parents=True, exist_ok=True)

    for c in df.columns:
        if df[c].dtype == "object":
            df[c] = df[c].astype("string")

    claim_parquet = data_dir / f"{output_stub}.parquet"
    df.to_parquet(claim_parquet, index=False)
    print(f"\n[save] {claim_parquet}")

    if claimant is not None:
        claimant_path = data_dir / f"{output_stub}_by_claimant.parquet"
        claimant.to_parquet(claimant_path, index=False)
        print(f"[save] {claimant_path}")

    if dealer is not None:
        dealer_path = data_dir / f"{output_stub}_by_dealer.parquet"
        dealer.to_parquet(dealer_path, index=False)
        print(f"[save] {dealer_path}")

    if save_excel:
        excel_path = data_dir / f"{output_stub}.xlsx"
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="claims_individual", index=False)
            if claimant is not None:
                claimant.to_excel(writer, sheet_name="by_claimant", index=False)
            if dealer is not None:
                dealer.to_excel(writer, sheet_name="by_dealer", index=False)

            # A small audit sheet makes the model inputs and policy settings explicit.
            settings = pd.DataFrame(
                {
                    "setting": [
                        "input_rows",
                        "model_feature_count",
                        "contamination",
                        "high_percentile",
                        "medium_percentile",
                        "wilson_z",
                        "medium_lift",
                        "medium_eb_probability",
                        "entity_counting_unit",
                    ],
                    "value": [
                        len(df),
                        len(used_features),
                        contamination,
                        HIGH_PERCENTILE,
                        MEDIUM_PERCENTILE,
                        WILSON_Z,
                        MEDIUM_LIFT,
                        MEDIUM_EB_PROB,
                        "unique ClaimId within entity when available; otherwise input row",
                    ],
                    "meaning": [
                        "Number of input records scored",
                        "Number of approved nonconstant numeric features used",
                        "Isolation Forest fit parameter; does not set exported High percentage",
                        "Global score percentile used for model High",
                        "Global score percentile used for model Medium",
                        "Confidence multiplier for Wilson lower bound",
                        "Configurable EB lift required for Medium",
                        "Configurable posterior probability required for Medium",
                        "Prevents duplicate claim-item rows from inflating entity statistics",
                    ],
                }
            )
            settings.to_excel(writer, sheet_name="model_settings", index=False)
            pd.DataFrame({"model_feature": used_features}).to_excel(
                writer, sheet_name="features_used", index=False
            )
            dropped.to_excel(writer, sheet_name="features_dropped", index=False)

        make_excel_readable(excel_path)
        print(
            f"[save] {excel_path} "
            "(claims_individual, by_claimant, by_dealer, model_settings, "
            "features_used, features_dropped)"
        )

    return df, claimant, dealer


def main():
    parser = argparse.ArgumentParser(
        description="Isolation Forest risk scoring for KM claims"
    )
    parser.add_argument("--input", default=INPUT_FILE)
    parser.add_argument("--output-stub", default=OUTPUT_STUB)
    parser.add_argument("--contamination", type=float, default=CONTAMINATION)
    parser.add_argument("--no-excel", action="store_true")
    args = parser.parse_args()

    run(
        input_path=args.input,
        output_stub=args.output_stub,
        contamination=args.contamination,
        save_excel=not args.no_excel,
    )


if __name__ == "__main__":
    main()
